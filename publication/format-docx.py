"""Apply print-only layout to Quarto's DOCX; never edit chapter sources."""
from pathlib import Path
import argparse, re
from urllib.parse import urlsplit, unquote
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

def value(el, name, val):
    child=el.find(qn(name))
    if child is None: child=OxmlElement(name); el.append(child)
    child.set(qn('w:val'),str(val)); return child

def format_book(source, target):
    d=Document(source); body=d.element.body
    original_math=len(body.xpath('.//m:oMath'))
    # Quarto's note boxes become single-column tables. Preserve every paragraph
    # and nested data table while removing only the decorative container/icon.
    for table in list(body.xpath('.//w:tbl'))[::-1]:
        rows=table.findall(qn('w:tr'))
        if not rows or any(len(row.findall(qn('w:tc')))!=1 for row in rows): continue
        for label in rows[0].xpath('.//w:p')[:1]:
            value(label.get_or_add_pPr(),'w:keepNext',1)
            for run in label.xpath('.//w:r'):
                value(run.get_or_add_rPr(),'w:b',1)
        parent=table.getparent(); pos=parent.index(table)
        for row in rows:
            cell=row.find(qn('w:tc'))
            for el in list(cell):
                if el.tag==qn('w:tcPr'):continue
                for drawing in el.xpath('.//w:drawing'):
                    drawing.getparent().remove(drawing)
                parent.insert(pos,el);pos+=1
        parent.remove(table)
    # Keep the approved homepage introduction/license, with a distinct print title.
    h1=[p for p in d.paragraphs if p.style.name=='Heading 1']
    assert len(h1)==20, len(h1)
    h1[0].text='版本与阅读说明'
    chapter_names=['index','01-introduction','03-work-environment','04-quick-start','method-learning','05-literature','theory-development','06-design','07-data-code','08-writing','02-project','10-automation','qualitative-research','11-safety-recovery','12-conclusion','appendix-a','appendix-b','appendix-c','references','ai-collaboration']
    chapter_bookmarks={name:f'v13_section_{i}' for i,name in enumerate(chapter_names)}
    bookmark_id=max(int(x.get(qn('w:id'))) for x in body.xpath('.//w:bookmarkStart'))+1
    for i,p in enumerate(h1):
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id+i));start.set(qn('w:name'),chapter_bookmarks[chapter_names[i]])
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id+i))
        p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)
    for number,p in enumerate(h1[2:14],1):
        run=p.add_run('')
        # Insert prefix without replacing the original heading/bookmark elements.
        r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=f'第{number}章  ';t.set(qn('xml:space'),'preserve');r.append(t)
        p._p.insert(1 if p._p.pPr is not None else 0,r)
        # Balance long chapter titles without changing their words or bookmarks.
        if number in {2,5,6,8,10}:
            for node in p._p.xpath('.//w:t'):
                if node.text and '：' in node.text:
                    left,right=node.text.split('：',1);node.text=left+'：'
                    br=OxmlElement('w:br');tail=OxmlElement('w:t');tail.text=right
                    node.addnext(br);br.addnext(tail);break
    subtitle='从零散问答到可推进、可检查、可交接的个人研究项目'
    for p in list(d.paragraphs):
        if (p.text==subtitle and p.style.name!='Subtitle') or p.text=='v1.3 · 2026-10-16':
            p._element.getparent().remove(p._element)
        elif p.style.name=='Date':p.text='v1.3 · 2026-10-16'
        elif not p.text.strip() and not p._p.xpath('.//m:oMath|.//w:drawing|.//w:bookmarkStart|.//w:fldChar|.//w:instrText'):
            p._element.getparent().remove(p._element)
    for section in d.sections:
        section.page_width=Cm(21);section.page_height=Cm(29.7)
        section.left_margin=Cm(3.175);section.right_margin=Cm(3.175)
        section.top_margin=Cm(2.54);section.bottom_margin=Cm(2.54)
        section.header_distance=Cm(1.27);section.footer_distance=Cm(1.27)
        section.different_first_page_header_footer=True
        footer=section.footer.paragraphs[0];footer.clear();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
        footer._p.append(field)
    for st in d.styles:
        if st.type!=1:continue
        st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        rpr=st.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
        for key in list(fonts.attrib):
            if 'Theme' in key or 'theme' in key:del fonts.attrib[key]
        fonts.set(qn('w:eastAsia'),'宋体')
    for name in ['Normal','Body Text','First Paragraph','Compact','Block Text']:
        st=d.styles[name];st.font.size=Pt(11.5)
        st.paragraph_format.line_spacing=1.3
        st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.widow_control=True
    for name,size in [('Title',28),('Subtitle',14),('Heading 1',17),('Heading 2',15),('Heading 3',12.5),('Heading 4',11.5)]:
        st=d.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle'
        st.element.rPr.find(qn('w:rFonts')).set(qn('w:eastAsia'),'微软雅黑')
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(16 if name.startswith('Heading') else 12)
        st.paragraph_format.space_after=Pt(8)
        st.paragraph_format.line_spacing=1.15
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    d.styles['Heading 1'].paragraph_format.page_break_before=True
    d.styles['Title'].paragraph_format.space_before=Pt(80)
    for name in ['Source Code','Verbatim Char']:
        if name in d.styles:
            d.styles[name].font.name='Consolas';d.styles[name].font.size=Pt(9.5)
    # Put the automatic contents on its own page; Word refreshes the field later.
    for sdt in body.xpath('./w:sdt'):
        paragraphs=sdt.findall('.//'+qn('w:p'))
        if paragraphs:
            ppr=paragraphs[0].get_or_add_pPr();value(ppr,'w:pageBreakBefore',1)
    width=d.sections[0].page_width-d.sections[0].left_margin-d.sections[0].right_margin
    for table in d.tables:
        n=len(table.columns);head=[cell.text for cell in table.rows[0].cells]
        if n==2: ratios=[.32,.68] if head[0] in ['材料','做法'] else [.42,.58]
        elif n==3:ratios=[.24,.38,.38]
        elif n==4:ratios=[.46,.18,.18,.18] if '规则' in head[0] else [.25]*4
        elif n==5:ratios=[.36,.16,.16,.16,.16] if '规则' in head[0] else [.18,.28,.18,.18,.18]
        else:ratios=[1/n]*n
        if n==2 and '编号' in head[0]:ratios=[.18,.82]
        if n==3 and head[0]=='片段':ratios=[.64,.18,.18]
        table.autofit=False
        value(table._tbl.tblPr,'w:tblLayout','fixed')
        borders=table._tbl.tblPr.find(qn('w:tblBorders'))
        if borders is None:borders=OxmlElement('w:tblBorders');table._tbl.tblPr.append(borders)
        for edge in ['top','left','bottom','right','insideH','insideV']:
            node=value(borders,'w:'+edge,'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9')
        for col,ratio in zip(table.columns,ratios):col.width=int(width*ratio)
        for i,row in enumerate(table.rows):
            pr=row._tr.get_or_add_trPr();value(pr,'w:cantSplit',1)
            if i==0:value(pr,'w:tblHeader',1)
            for j,cell in enumerate(row.cells):
                cell.width=int(width*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    node=OxmlElement('w:'+side);node.set(qn('w:w'),'85');node.set(qn('w:type'),'dxa');mar.append(node)
                old=cp.find(qn('w:tcMar'))
                if old is not None:cp.remove(old)
                cp.append(mar)
                shd=cp.find(qn('w:shd'))
                if shd is None:shd=OxmlElement('w:shd');cp.append(shd)
                shd.set(qn('w:fill'),'EAF0F5' if i==0 else 'FFFFFF')
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.keep_with_next=i==0
                    for run in p.runs:
                        run.font.size=Pt(10.5)
                        if i==0:run.font.bold=True
        # Preserve row content while allowing a long table to break between rows.
    for shape in d.inline_shapes:
        if shape.width>width:
            factor=width/shape.width;shape.width=int(width);shape.height=int(shape.height*factor)
    for p in d.paragraphs:
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    assert len(body.xpath('.//m:oMath'))==original_math
    # Every active external link must be a real URL; unused Pandoc rels are harmless.
    for link in body.xpath('.//w:hyperlink'):
        rid=link.get(qn('r:id'))
        if rid:
            rel=d.part.rels[rid]
            url=urlsplit(str(rel.target_ref))
            if not url.scheme and url.path.endswith('.qmd'):
                fragment=unquote(url.fragment) or link.get(qn('w:anchor'))
                anchor=fragment or chapter_bookmarks[Path(url.path).stem]
                assert anchor in {x.get(qn('w:name')) for x in body.xpath('.//w:bookmarkStart')},anchor
                link.set(qn('w:anchor'),anchor);del link.attrib[qn('r:id')]
                continue
            assert str(rel.target_ref).startswith(('https://','http://','mailto:')),rel.target_ref
    d.core_properties.title='心理学研究的 AI 进阶手册'
    d.core_properties.author='Hao Xie';d.core_properties.last_modified_by='Hao Xie'
    d.core_properties.subject='v1.3 · 2026-10-16'
    target.parent.mkdir(parents=True,exist_ok=True);d.save(target)
    print(f'Formatted {target.name}: {len(d.tables)} data tables, {len(d.inline_shapes)} figures, {original_math} native equations.')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('input',type=Path);p.add_argument('output',type=Path);a=p.parse_args();format_book(a.input,a.output)
